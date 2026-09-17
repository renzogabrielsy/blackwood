import sys, collections
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as L
A = load_workbook("original.xlsx"); B = load_workbook("renzo-edited.xlsx")
def col(c):
    if c is None: return None
    return c.rgb if c.type == "rgb" else (f"theme{c.theme}" + (f"/{c.tint:+.2f}" if c.tint else "")) if c.type == "theme" else f"idx{c.indexed}"
def sty(c):
    f, fl, b, a = c.font, c.fill, c.border, c.alignment
    side = lambda s: (s.style, col(s.color)) if s and s.style else None
    return dict(font=(f.name, f.sz, f.b, f.i, col(f.color)), fill=(fl.fill_type, col(fl.fgColor) if fl.fill_type else None),
                border=(side(b.left), side(b.right), side(b.top), side(b.bottom)), align=(a.horizontal, a.vertical, a.wrap_text, a.indent),
                fmt=c.number_format)
for name in B.sheetnames:
    a, b = A[name], B[name]
    out = []
    # sheet-level
    for k, fa, fb in [("freeze", a.freeze_panes, b.freeze_panes), ("gridlines", a.sheet_view.showGridLines, b.sheet_view.showGridLines), ("zoom", a.sheet_view.zoomScale, b.sheet_view.zoomScale),
                      ("tabColor", col(a.sheet_properties.tabColor), col(b.sheet_properties.tabColor)), ("dims", a.dimensions, b.dimensions),
                      ("merged", sorted(map(str, a.merged_cells.ranges)), sorted(map(str, b.merged_cells.ranges))),
                      ("cond_fmt", len(a.conditional_formatting), len(b.conditional_formatting)), ("orientation", a.page_setup.orientation, b.page_setup.orientation)]:
        if fa != fb: out.append(f"  SHEET {k}: {fa} -> {fb}")
    cw = {k: v.width for k, v in a.column_dimensions.items()}; cwb = {k: (v.width, v.hidden) for k, v in b.column_dimensions.items()}
    for k in sorted(set(cw) | set(cwb), key=lambda x: (len(x), x)):
        wa = cw.get(k); wbv = cwb.get(k, (None, False))
        if wa != wbv[0] or wbv[1]: out.append(f"  COL {k}: width {wa} -> {wbv[0]}{' HIDDEN' if wbv[1] else ''}")
    rh = {k: v.height for k, v in a.row_dimensions.items() if v.height}; rhb = {k: (v.height, v.hidden) for k, v in b.row_dimensions.items() if v.height or v.hidden}
    rows_changed = collections.Counter()
    for k in sorted(set(rh) | set(rhb)):
        ha = rh.get(k); hb = rhb.get(k, (None, False))
        if ha != hb[0] or hb[1]: rows_changed[(ha, hb[0], hb[1])] += 1; 
    for (ha, hb, hid), n in rows_changed.items(): out.append(f"  ROW HEIGHT {ha} -> {hb}{' HIDDEN' if hid else ''}  x{n} rows")
    # cells
    groups = collections.defaultdict(list); vals = []
    maxr = max(a.max_row, b.max_row); maxc = max(a.max_column, b.max_column)
    for r in range(1, maxr + 1):
        for c in range(1, maxc + 1):
            ca, cb = a.cell(r, c), b.cell(r, c)
            if ca.value != cb.value: vals.append((ca.coordinate, ca.value, cb.value))
            sa, sb = sty(ca), sty(cb)
            for k in sa:
                if sa[k] != sb[k]: groups[(k, str(sa[k]), str(sb[k]))].append(cb.coordinate)
    for (k, va, vb), cells in sorted(groups.items(), key=lambda x: -len(x[1])):
        rows = sorted({int(''.join(ch for ch in c if ch.isdigit())) for c in cells}); 
        out.append(f"  STYLE {k}: {va} -> {vb}   [{len(cells)} cells, rows {rows[0]}-{rows[-1]}, e.g. {', '.join(cells[:6])}]")
    for v in vals[:25]: out.append(f"  VALUE {v[0]}: {v[1]!r} -> {v[2]!r}")
    if len(vals) > 25: out.append(f"  ... {len(vals)-25} more value changes")
    print(f"=== {name}: {'NO CHANGES' if not out else str(len(out)) + ' change groups'}"); print("\n".join(out))
